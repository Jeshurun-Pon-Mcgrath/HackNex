from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import Settings
from app.main import create_app


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        database_path=tmp_path / "metadata.sqlite3",
        allowed_origins=["http://localhost:5173"],
        max_upload_bytes=1024 * 1024,
    )


@pytest.fixture
def client(settings: Settings) -> Iterator[TestClient]:
    with TestClient(create_app(settings)) as test_client:
        yield test_client


@pytest.fixture
def csv_bytes() -> bytes:
    return (
        b"region,profit,cost,date,price\n"
        b"North,10,2,01/02/2024,$10\n"
        b"North,10,2,01/02/2024,EUR 10\n"
        b"South,100,0,13/02/2024,$20\n"
    )


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> bytes:
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Orders"
    first.append(["region", "amount"])
    first.append(["North", 10])
    second = workbook.create_sheet("Returns")
    second.append(["region", "amount"])
    second.append(["South", 2])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["secret"])
    hidden.append(["value"])
    workbook.save(path)
    return path.read_bytes()
