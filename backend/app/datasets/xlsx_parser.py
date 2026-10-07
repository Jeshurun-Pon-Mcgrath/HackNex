from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import AppError
from app.datasets.csv_parser import ParsedTable
from app.datasets.models import WorksheetInfo


def inspect_workbook(path: Path) -> list[WorksheetInfo]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise AppError(422, "malformed_workbook", "The workbook could not be read safely.") from exc
    try:
        result: list[WorksheetInfo] = []
        for index, sheet in enumerate(workbook.worksheets):
            usable = any(
                any(value is not None for value in row) for row in sheet.iter_rows(values_only=True)
            )
            result.append(
                WorksheetInfo(
                    name=sheet.title,
                    index=index,
                    visibility=sheet.sheet_state,
                    usable=usable,
                )
            )
        return result
    finally:
        workbook.close()


def parse_sheet(path: Path, sheet_name: str) -> ParsedTable:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise AppError(422, "unknown_sheet", "The selected worksheet does not exist.")
        sheet = workbook[sheet_name]
        if sheet.sheet_state != "visible":
            raise AppError(422, "hidden_sheet", "Hidden worksheets cannot be selected.")
        iterator = sheet.iter_rows(values_only=False)
        header_cells = next(
            (row for row in iterator if any(cell.value is not None for cell in row)), None
        )
        if header_cells is None:
            raise AppError(422, "no_usable_header", "The worksheet is empty.")
        headers = ["" if cell.value is None else str(cell.value) for cell in header_cells]
        rows: list[list[Any]] = []
        formula_columns: set[int] = set()
        for row in iterator:
            row_values: list[Any] = []
            for index, cell in enumerate(row[: len(headers)]):
                if cell.data_type == "f":
                    formula_columns.add(index)
                    row_values.append(None)
                else:
                    row_values.append(cell.value)
            if any(
                value is not None and (not isinstance(value, str) or value.strip())
                for value in row_values
            ):
                rows.append(row_values)
        if not headers or all(not header.strip() for header in headers):
            raise AppError(422, "no_usable_header", "The worksheet has no usable header.")
        if not rows:
            raise AppError(422, "no_data_rows", "The worksheet has no usable data rows.")
        warnings = (
            ["Formula cells were detected and were not evaluated."] if formula_columns else []
        )
        return ParsedTable(headers, rows, warnings, formula_columns)
    finally:
        workbook.close()
